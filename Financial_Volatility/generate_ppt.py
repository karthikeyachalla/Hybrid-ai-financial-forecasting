import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_dynamic_ppt():
    prs = Presentation()
    
    # Set slide width and height to 16:9 ratio (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Theme Colors
    BG_COLOR = RGBColor(18, 24, 38)        # Deep Dark Blue
    TEXT_COLOR = RGBColor(240, 245, 250)   # Off White for readability
    ACCENT_COLOR = RGBColor(0, 210, 255)   # Neon Cyan
    SECONDARY_ACCENT = RGBColor(128, 90, 213) # Purple accent for variety

    def apply_background_and_accents(slide):
        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        
        # Left Accent Bar
        sidebar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), prs.slide_height
        )
        sidebar.fill.solid()
        sidebar.fill.fore_color.rgb = ACCENT_COLOR
        sidebar.line.fill.background()
        
        # Top right decorative shape
        top_dec = slide.shapes.add_shape(
            MSO_SHAPE.PARALLELOGRAM, Inches(11.5), Inches(-0.2), Inches(2.5), Inches(0.6)
        )
        top_dec.fill.solid()
        top_dec.fill.fore_color.rgb = SECONDARY_ACCENT
        top_dec.line.fill.background()

    def format_text(shape, font_size, color=TEXT_COLOR, bold=False, align=None, is_title=False):
        for paragraph in shape.text_frame.paragraphs:
            if align:
                paragraph.alignment = align
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
                run.font.color.rgb = color
                run.font.bold = bold
                if is_title:
                    run.font.name = 'Arial'

    def add_title_slide(title_text, subtitle_text):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
        apply_background_and_accents(slide)
        
        # Add a center floating box
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2), Inches(10.333), Inches(3.5)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(25, 35, 55)
        box.line.color.rgb = ACCENT_COLOR
        box.line.width = Pt(2)
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9.333), Inches(1.5))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.text = title_text
        format_text(title_box, 60, color=ACCENT_COLOR, bold=True, align=PP_ALIGN.CENTER, is_title=True)
        
        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(2), Inches(4), Inches(9.333), Inches(1.5))
        sf = sub_box.text_frame
        sf.word_wrap = True
        sf.text = subtitle_text
        format_text(sub_box, 26, color=TEXT_COLOR, align=PP_ALIGN.CENTER)

    def add_content_slide(title_text, bullet_points, image_path=None):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
        apply_background_and_accents(slide)
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        tf = title_box.text_frame
        tf.text = title_text
        format_text(title_box, 40, color=ACCENT_COLOR, bold=True, is_title=True)
        
        # If an image exists, we split the slide. 6 inches for text, rest for image.
        has_image = image_path and os.path.exists(image_path)
        content_width = Inches(5.8) if has_image else Inches(11.5)
        
        # Content Box
        body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), content_width, Inches(5.2))
        bf = body_box.text_frame
        bf.word_wrap = True
        
        for i, point in enumerate(bullet_points):
            p = bf.add_paragraph() if i > 0 else bf.paragraphs[0]
            p.text = "✦  " + point
            p.level = 0
            p.space_after = Pt(14)
            p.space_before = Pt(6)
            
        format_text(body_box, 24)

        # Image formatting
        if has_image:
            # We want to fit the image organically on the right side
            # Slide size is 13.333 x 7.5. Right side starts around 6.5
            pic = slide.shapes.add_picture(image_path, Inches(6.8), Inches(1.6), width=Inches(6.0))
            # Add a stylish frame behind the image
            frame = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.7), Inches(6.0), pic.height
            )
            frame.fill.solid()
            frame.fill.fore_color.rgb = SECONDARY_ACCENT
            frame.line.fill.background()
            # Push frame to back (by creating it after, pic will be underneath, so let's swap z-order logic)
            # Actually, we can just position the pic again or adjust the code. To keep it simple, let's load frame first.

    def add_content_slide_with_frame(title_text, bullet_points, image_path=None):
        slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
        apply_background_and_accents(slide)
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        tf = title_box.text_frame
        tf.text = title_text
        format_text(title_box, 40, color=ACCENT_COLOR, bold=True, is_title=True)
        
        has_image = image_path and os.path.exists(image_path)
        content_width = Inches(5.8) if has_image else Inches(11.5)
        
        # Frame and Pic
        if has_image:
            # Add frame first so it goes behind
            frame = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(7.0), Inches(1.8), Inches(5.8), Inches(4.5)
            )
            frame.fill.solid()
            frame.fill.fore_color.rgb = SECONDARY_ACCENT
            frame.line.fill.background()
            
            # Then image
            slide.shapes.add_picture(image_path, Inches(6.8), Inches(1.6), width=Inches(5.8))
        
        # Content Box
        body_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), content_width, Inches(5.2))
        bf = body_box.text_frame
        bf.word_wrap = True
        
        for i, point in enumerate(bullet_points):
            p = bf.add_paragraph() if i > 0 else bf.paragraphs[0]
            p.text = "✦  " + point
            p.level = 0
            p.space_after = Pt(14)
            p.space_before = Pt(6)
            
        format_text(body_box, 24)


    # --- SLIDES GENERATION ---

    # 1. Title Slide
    add_title_slide(
        "Hybrid AI Financial Forecasting", 
        "Volatility & Market Regimes\nPowered by Deep Learning, VAE, and NLP\nQUANT-PRO v4.1"
    )

    # 2. Challenge
    add_content_slide_with_frame(
        "1. The Challenge with Traditional Models",
        [
            "Financial markets are inherently noisy, non-stationary, and highly complex.",
            "Classical statistical models rely strictly on linear assumptions.",
            "They fail to adapt dynamically to sudden regime shifts & extreme volatility crashes.",
            "Purely mathematical models ignore human sentiment and unstructured qualitative data (News, Twitter)."
        ],
        "assets/returns_plot.png"
    )

    # 3. Architecture
    add_content_slide_with_frame(
        "2. The Solution: QUANT-PRO Architecture",
        [
            "A hybrid quantitative framework blending classical stats with state-of-the-art Deep Learning.",
            "End-to-End automated data pipeline starting from live market extraction.",
            "Core components: Non-linear Volatility Modeling, Anomaly Scorer, Sentiment Engine, Price Forecaster.",
            "Designed for intuitive UI/UX with clear risk metrics & neural explainability."
        ]
    )

    # 4. Feature Engineering
    add_content_slide_with_frame(
        "3. Live Data & Feature Engineering",
        [
            "Extracts real-time price & volume streams seamlessly.",
            "Dynamic automated feature engineering computes technical indicators instantly.",
            "Key Alpha Generators:",
            "  • Relative Strength Index (RSI)",
            "  • Moving Average Convergence Divergence (MACD)",
            "  • Bollinger Bands & Momentum Oscillators",
            "Creates a robust multi-dimensional feature set bridging price action & momentum."
        ],
        "assets/acf_plots.png"
    )

    # 5. Volatility Engine
    add_content_slide_with_frame(
        "4. Volatility Engine: GARCH + LSTM",
        [
            "Uses Realized GARCH to capture baseline mean-reverting behavior.",
            "Implements a bidirectional LSTM network for capturing residual non-linear chaos.",
            "The Hybrid Blend:",
            "  • Extracts structural variance via econometrics.",
            "  • Learns hidden, non-linear market regimes utilizing deep networks.",
            "  • Identifies the accurate probability of future market stress."
        ],
        "assets/volatility_models_comparison.png"
    )

    # 6. Regime Detection
    add_content_slide_with_frame(
        "5. Market Regime Detection (VAE)",
        [
            "Markets endlessly transition between 'Bull', 'Bear', and 'Sideways' modes.",
            "We utilize a Variational Autoencoder (VAE) for unsupervised anomaly scoring.",
            "Continuously reconstructs recent sequences and computes a structural Reconstruction Loss.",
            "Elevated anomalies signal market stress ('Volatile' or 'Panic').",
            "Low scores confirm structural integrity ('Normal'/'Steady')."
        ],
        "assets/market_regimes.png"
    )

    # 7. Sentiment
    add_content_slide_with_frame(
        "6. Sentiment Language Engine",
        [
            "Extracts live psychological factors utilizing HuggingFace's FinBERT model.",
            "Automatically pulls top financial news headlines specific to the tracked asset.",
            "Evaluates localized context, trained exclusively on Wall Street lexicon.",
            "Computes live probabilistic Bullish vs Bearish scoring in real-time.",
            "Mitigates the blind-spots of pure mathematical chart models."
        ]
    )

    # 8. Forecasting
    add_content_slide_with_frame(
        "7. Price Forecasting & 'Fast Mode'",
        [
            "Employs Multi-Horizon LSTMs to predict localized price trajectories.",
            "Quantile Regression maps absolute confidence intervals (10th/90th percentile).",
            "The 'Fast Mode' edge:",
            "  • Pre-trained globally on immense historical datasets.",
            "  • Pre-compiled neural weights auto-load locally.",
            "  • Accelerates local rendering speeds by 95% while drastically reducing compute requirements."
        ],
        "assets/confidence_intervals.png"
    )

    # 9. Conclusion
    add_content_slide_with_frame(
        "8. Conclusion & Core Impact",
        [
            "Scales hedge-fund tier quantitative metrics into accessible, user-centric dashboards.",
            "Melds traditional econometrics with modern Deep Learning and NLP extraction.",
            "Empowers transparent analysis via Explainable AI architectures.",
            "Rapidly plug-and-playable for Equities, Crypto, and Global Forex architectures."
        ]
    )

    os.makedirs("docs", exist_ok=True)
    output_path = "docs/QuantVision_Presentation.pptx"
    prs.save(output_path)
    print(f"Successfully generated dynamic 9-slide presentation at {output_path}")

if __name__ == "__main__":
    create_dynamic_ppt()
